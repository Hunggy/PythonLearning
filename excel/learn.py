import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
# 任务：新建工作簿 -> 拿到默认工作表 -> A1 写入"学生姓名" -> 保存为「成绩表.xlsx」
# 提示：wb = openpyxl.Workbook() / ws = wb.active / ws["A1"] = "学生姓名" / wb.save("成绩表.xlsx")


wb = openpyxl.Workbook()
ws = wb.active

headers = ["学生姓名", "语文", "数学", "英语"]
for col, header in enumerate(headers, start=1):
    ws.cell(row=1, column=col, value=header)


scores = [["张三", 85, 92, 78], ["李四", 90, 88, 95], ["王五", 76, 84, 89]]
while True:
    add_more = input("是否继续添加学生成绩？(y/n)：")
    if add_more.lower() != 'y':
        break
    scores.append(input("请输入学生姓名和成绩（格式：姓名 语文 数学 英语）：").split())
for row_num, row in enumerate(scores, start=2):
    for col, value in enumerate(row, start=1):
        if col > 1:   # 姓名保留字符串，成绩转 int
            value = int(value)
        ws.cell(row=row_num, column=col, value=value)

for col in range(1, 5):   # 1~4 列
    cell = ws.cell(row=1, column=col)   # 拿到单元格
    cell.font = Font(bold=True)         # 加粗
    cell.alignment = Alignment(horizontal="center")  # 居中

for col in range(1, 5):   # 1~4 列
    ws.column_dimensions[get_column_letter(col)].width = 8

wb.save("成绩表.xlsx")